"""Extracción de texto de PDF/DOCX/TXT y señales de formato para el ATS.

Además del texto, devolvemos `DocumentSignals`: cosas que un ATS real penaliza y
que solo se pueden medir mirando el archivo, no el texto ya extraído (tablas,
imágenes, multi-columna, fuentes raras).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import pdfplumber
from docx import Document as DocxDocument

from app.core.logging import get_logger

log = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class DocumentSignals:
    pages: int = 0
    words: int = 0
    tables: int = 0
    images: int = 0
    distinct_fonts: set[str] = field(default_factory=set)
    likely_multicolumn: bool = False
    has_text_layer: bool = True
    bullet_chars: int = 0

    def as_dict(self) -> dict:
        return {
            "pages": self.pages,
            "words": self.words,
            "tables": self.tables,
            "images": self.images,
            "distinct_fonts": sorted(self.distinct_fonts),
            "likely_multicolumn": self.likely_multicolumn,
            "has_text_layer": self.has_text_layer,
            "bullet_chars": self.bullet_chars,
        }


@dataclass
class ParsedDocument:
    text: str
    signals: DocumentSignals


class UnsupportedDocument(ValueError):
    pass


def parse_document(data: bytes, filename: str) -> ParsedDocument:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == ".pdf":
        return _parse_pdf(data)
    if ext == ".docx":
        return _parse_docx(data)
    if ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
        return ParsedDocument(text=clean_text(text), signals=_text_signals(text))
    if ext == ".doc":
        raise UnsupportedDocument(
            "El formato .doc (Word 97) no se soporta. Guárdalo como .docx o PDF."
        )
    raise UnsupportedDocument(
        f"Extensión no soportada: '{ext or 'sin extensión'}'. "
        f"Usa: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _parse_pdf(data: bytes) -> ParsedDocument:
    signals = DocumentSignals()
    chunks: list[str] = []

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        signals.pages = len(pdf.pages)
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            chunks.append(page_text)

            signals.tables += len(page.find_tables())
            signals.images += len(page.images)
            for char in page.chars:
                font = char.get("fontname")
                if font:
                    # "ABCDEF+Arial-Bold" -> "Arial"
                    signals.distinct_fonts.add(re.sub(r"^[A-Z]{6}\+|[-,].*$", "", font))

            if _looks_multicolumn(page):
                signals.likely_multicolumn = True

    text = "\n".join(chunks)
    signals.has_text_layer = len(text.strip()) >= 40
    if not signals.has_text_layer:
        log.warning("El PDF no tiene capa de texto — probablemente es un escaneo.")
    signals.words = len(text.split())
    signals.bullet_chars = len(re.findall(r"[•▪●◦‣·]", text))
    return ParsedDocument(text=clean_text(text), signals=signals)


def _looks_multicolumn(page) -> bool:
    """Heurística: si hay dos bloques densos de palabras separados por un pasillo
    vertical vacío en el centro, el PDF es de dos columnas — y muchos ATS leen
    ese layout en el orden equivocado."""
    words = page.extract_words() or []
    if len(words) < 60:
        return False
    width = page.width or 1
    left = sum(1 for w in words if w["x1"] < width * 0.45)
    right = sum(1 for w in words if w["x0"] > width * 0.55)
    middle = sum(1 for w in words if w["x0"] <= width * 0.55 and w["x1"] >= width * 0.45)
    return left > 25 and right > 25 and middle < len(words) * 0.12


def _parse_docx(data: bytes) -> ParsedDocument:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]

    tables = 0
    for table in doc.tables:
        tables += 1
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    text = "\n".join(parts)
    signals = _text_signals(text)
    signals.tables = tables
    signals.images = len(doc.inline_shapes)
    signals.distinct_fonts = {
        r.font.name for p in doc.paragraphs for r in p.runs if r.font and r.font.name
    }
    # Word no pagina hasta renderizar; ~500 palabras por página es una aproximación útil.
    signals.pages = max(1, round(signals.words / 500))
    return ParsedDocument(text=clean_text(text), signals=signals)


def _text_signals(text: str) -> DocumentSignals:
    return DocumentSignals(
        pages=max(1, round(len(text.split()) / 500)),
        words=len(text.split()),
        bullet_chars=len(re.findall(r"[•▪●◦‣·]", text)),
    )


def clean_text(text: str) -> str:
    """Normaliza el texto extraído sin destruir la estructura de líneas."""
    text = text.replace(" ", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return text.strip()


def html_to_text(html: str) -> str:
    """Convierte la descripción HTML de una oferta a texto plano legible."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for li in soup.find_all("li"):
        li.insert(0, "• ")
    return clean_text(soup.get_text("\n"))
